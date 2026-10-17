"""Единый движок документа: HTML-координаты или Word-шаблон (DOCX → PDF).

Один вход — `build_document(code, values, out_dir, ...)` — и система сама выбирает путь
по полю `templates.engine`: 'html' (координатная сетка, weasyprint) или 'word'
(DOCX-шаблон + LibreOffice). Так документ собирается одним способом и там, где его
смотрит конструктор, и там, где его выпускает заявка.
"""
from __future__ import annotations

import os
import re
import shutil
import subprocess
import zipfile
from pathlib import Path

ROOT = Path("/opt/wolf_v2")
TPL_ROOT = ROOT / "data" / "templates"
SOFFICE = shutil.which("soffice") or "/usr/bin/soffice"

_missing_cache: dict[str, list] = {}


def engine_of(code: str, engine: str | None = None) -> str:
    """Движок документа: из базы, если не передан явно. По умолчанию html."""
    if engine:
        return engine
    try:
        import sqlite3
        con = sqlite3.connect(str(ROOT / "data" / "wolf.db"))
        row = con.execute("SELECT engine FROM templates WHERE code=?", (code,)).fetchone()
        con.close()
        if row and row[0]:
            return row[0]
    except Exception:
        pass
    return "html"


def word_template(code: str) -> Path:
    return TPL_ROOT / code / "word" / "template.docx"


def docx_variables(code: str) -> list:
    """Имена переменных внутри DOCX-шаблона (jinja может рвать {{ x }} между run'ами)."""
    path = word_template(code)
    if not path.exists():
        return []
    try:
        with zipfile.ZipFile(str(path)) as z:
            xml = z.read("word/document.xml").decode("utf-8", "ignore")
    except Exception:
        return []
    xml = re.sub(r"<[^>]+>", "", xml)
    return sorted(set(re.findall(r"\{\{\s*(\w+)\s*\}\}", xml)))


def placeholders_left(path: Path) -> list:
    """Остались ли в PDF незаполненные метки {{...}} — документ выпускать нельзя."""
    try:
        import pymupdf
        d = pymupdf.open(str(path))
        text = " ".join(p.get_text() for p in d)
        d.close()
    except Exception:
        return []
    return sorted(set(re.findall(r"\{\{[^}]{0,40}\}\}", text)))


_ARIAL_BOLD = "/usr/share/fonts/truetype/arial/Arial-Bold.ttf"
_ARIAL = "/usr/share/fonts/truetype/arial/Arial.ttf"
_font_cache: dict = {}


def _text_width(text: str, size: float, bold: bool = True) -> float:
    """Ширина строки в pt шрифтом значений документа (Arial / Arial Bold)."""
    from PIL import ImageFont
    path = _ARIAL_BOLD if bold else _ARIAL
    f = _font_cache.get(path)
    if f is None:
        f = ImageFont.truetype(path, 64)
        _font_cache[path] = f
    return f.getlength(text) / 64.0 * size


def _wrap_lines(text: str, size: float, w_pt: float, bold: bool = True) -> int:
    """Сколько строк займёт значение в рамке шириной w_pt (как переносит LibreOffice/Word)."""
    lines, cur = 1, ""
    for word in str(text).split():
        probe = (cur + " " + word).strip()
        if not cur or _text_width(probe, size, bold) <= w_pt:
            cur = probe
        else:
            lines += 1
            cur = word
    return lines


def _fit_size(value, w_pt: float, max_lines: int = 1, max_height: float = 19.0,
              base: float = 9.5, bold: bool = True, min_size: float = 5.0,
              width_safety: float = 0.97) -> float:
    """Наибольший кегль, при котором значение укладывается в max_lines строк.

    Ширина берётся с запасом (width_safety): шрифтометрия PIL и рендер LibreOffice
    расходятся на десятые доли pt. Высота строк (1.18*кегль) не должна превышать
    max_height — иначе многострочное значение сядет на следующую строку полей.
    """
    if not value:
        return base
    size = base
    while size > min_size:
        wlim = w_pt * width_safety
        lines = _wrap_lines(value, size, wlim, bold)
        if lines <= max_lines and lines * 1.18 * size <= max_height:
            return round(size, 1)
        size -= 0.5
    return min_size


def fit_word_template(tpl: Path, vals: dict, out: Path, rules: dict | None = None) -> Path:
    """Копия DOCX-шаблона с кеглями, подогнанными под длину значений.

    rules: {переменная: макс_строк} или {переменная: (макс_строк, макс_высота_pt)};
    по умолчанию 1 строка. Ширина берётся из рамки абзаца (w:framePr/w:w) — то есть
    ограничение задано в шаблоне и живёт вместе с ним.
    """
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn
    rules = dict(rules or {})
    doc = Document(str(tpl))
    for p in doc.paragraphs:
        m = re.fullmatch(r"\{\{\s*([A-Za-z_]\w*)\s*\}\}", (p.text or "").strip())
        if not m:
            continue
        name = m.group(1)
        pPr = p._p.find(qn("w:pPr"))
        fr = pPr.find(qn("w:framePr")) if pPr is not None else None
        w_pt = 300.0
        if fr is not None and fr.get(qn("w:w")):
            w_pt = int(fr.get(qn("w:w"))) / 20.0
        if name == "address_in_indonesia":
            # адрес — строго одна строка: сокращаем текст, ширину рамки расширяем до края листа
            from services.address_fit import SHEET_RIGHT, VALUE_X0, fit_address
            txt, size, _short = fit_address(str(vals.get(name) or ""))
            vals[name] = txt
            if fr is not None:
                # рамка шире строки, чтобы Wert не переносился: адрес — всегда одна строка
                fr.set(qn("w:w"), str(int((SHEET_RIGHT - VALUE_X0) * 20)))
            for r in p.runs:
                r.font.size = Pt(size)
            continue
        rule = rules.get(name, 1)
        if isinstance(rule, (tuple, list)):
            max_lines, max_h = int(rule[0]), float(rule[1])
        else:
            max_lines, max_h = int(rule), 19.0
        size = _fit_size(vals.get(name), w_pt, max_lines, max_h)
        for r in p.runs:
            r.font.size = Pt(size)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def build_word(code: str, values: dict, out_dir: Path | str, name: str | None = None,
               timeout: int = 240) -> tuple[Path | None, Path | None, list]:
    """Собрать документ из Word-шаблона. Возвращает (docx, pdf, пустые_переменные)."""
    tpl = word_template(code)
    if not tpl.exists():
        return None, None, []
    from docxtpl import DocxTemplate

    vals = {k: ("" if v is None else str(v)) for k, v in (values or {}).items()}
    missing = [v for v in docx_variables(code) if not vals.get(v)]
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^\w.-]+", "_", (name or code))[:80]
    out_docx = out_dir / ("%s.docx" % safe)
    doc = DocxTemplate(str(fit_word_template(tpl, vals, out_dir / (".fit_%s.docx" % safe),
                                             rules={})))
    doc.render(vals)
    doc.save(str(out_docx))
    subprocess.run([SOFFICE, "--headless", "--norestore", "--convert-to", "pdf",
                    "--outdir", str(out_dir), str(out_docx)],
                   capture_output=True, timeout=timeout)
    out_pdf = out_docx.with_suffix(".pdf")
    return out_docx, (out_pdf if out_pdf.exists() else None), missing


def build_html(code: str, values: dict, out_pdf: Path | str, keep_unfilled: bool = True):
    """Собрать документ из HTML-шаблона (координатная сетка). Возвращает (путь | None)."""
    from services import render as doc_render
    html_file = TPL_ROOT / code / "template.html"
    if not html_file.exists():
        return None
    html = html_file.read_text(encoding="utf-8")
    filled = doc_render.fill(html, values or {}, keep_unfilled=keep_unfilled, template_code=code)
    out_pdf = Path(out_pdf)
    out_pdf.parent.mkdir(parents=True, exist_ok=True)
    doc_render.render_pdf(filled, out_path=str(out_pdf))
    return out_pdf if out_pdf.exists() else None


def build_document(code: str, values: dict, out_dir: Path | str, engine: str | None = None,
                   name: str | None = None, out_pdf: Path | str | None = None) -> dict:
    """Собрать документ тем движком, который закреплён за шаблоном.

    Возвращает {'ok', 'engine', 'pdf', 'docx', 'missing', 'placeholders'}.
    """
    eng = engine_of(code, engine)
    out_dir = Path(out_dir)
    res = {"ok": False, "engine": eng, "pdf": None, "docx": None, "missing": [], "placeholders": []}
    if eng == "word":
        docx, pdf, missing = build_word(code, values, out_dir, name=name)
        res.update(docx=str(docx) if docx else None, pdf=str(pdf) if pdf else None, missing=missing)
    else:
        target = Path(out_pdf) if out_pdf else (out_dir / ("%s.pdf" % re.sub(r"[^\w.-]+", "_", (name or code))[:80]))
        pdf = build_html(code, values, target)
        res.update(pdf=str(pdf) if pdf else None)
    if res["pdf"]:
        res["placeholders"] = placeholders_left(Path(res["pdf"]))
        res["ok"] = True
    return res
