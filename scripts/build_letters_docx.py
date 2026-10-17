#!/usr/bin/env python3
"""Единый каркас писем иммиграции: Bridging (№9) и Conversion (№10).

Оба письма — один и тот же бланк PT ALPHA PRIME BALI: шапка, Nomor/Lampiran/Hal,
адресат, блок подписанта, блок клиента, подпись. Отличаются только значения
(тема письма и строка о предмете обращения приходят из заявки).

Запуск: /usr/bin/python3 scripts/build_letters_docx.py [--render]
"""
from __future__ import annotations

import re
import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path

import docxtpl
import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path('/opt/wolf_v2')
CODES = ('bridging_request_letter__default', 'conversion_request_letter__default')
LETTERHEAD = [
    'PT ALPHA PRIME BALI',
    'Jl. Wisma Udayana No.164 Buruan, Kec. Blahbatuh, Kabupaten Gianyar, Bali 80581',
    'Telp: +62 811 397 9669, Email: ptalphaprimeconsulting@gmail.com',
]


def p(doc, text='', size=11, bold=False, align=None, before=0, after=0):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), 1.15
    if align is not None:
        par.alignment = align
    if text:
        r = par.add_run(text)
        r.font.size, r.bold, r.font.name = Pt(size), bold, 'Times New Roman'
    return par


def field(doc, label, value, size=11):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = 1.15
    pf.tab_stops.add_tab_stop(Cm(4.2))
    pf.tab_stops.add_tab_stop(Cm(4.6))
    r = par.add_run(f'{label}\t:\t{value}')
    r.font.size, r.font.name = Pt(size), 'Times New Roman'


def build(code: str) -> Path:
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(1.8)
    s.left_margin, s.right_margin = Cm(2.2), Cm(2.0)
    doc.styles['Normal'].font.name = 'Times New Roman'
    doc.styles['Normal'].font.size = Pt(11)

    p(doc, LETTERHEAD[0], size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    p(doc, LETTERHEAD[1], size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    p(doc, LETTERHEAD[2], size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    line = doc.add_paragraph()
    line.paragraph_format.space_after = Pt(8)
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '000000')
    pbdr.append(bottom)
    line._p.get_or_add_pPr().append(pbdr)

    p(doc, 'Nomor: {{ letter_no }}')
    p(doc, 'Lampiran: satu berkas')
    p(doc, 'Hal: {{ letter_subject }}', after=10)
    p(doc, 'Kepada Yth.')
    p(doc, '{{ addressee_line1 }}')
    p(doc, '{{ addressee_line2 }}')
    p(doc, '{{ addressee_line3 }}', after=10)

    p(doc, 'Bersama ini dengan hormat kami:', after=2)
    field(doc, 'Nama', '{{ attorney_name }}')
    field(doc, 'Pekerjaan', 'Swasta')
    field(doc, 'Kewarganegaraan', '{{ attorney_citizenship }}')
    field(doc, 'NIK', '{{ attorney_nik }}')
    field(doc, 'Alamat', '{{ attorney_address }}')

    p(doc, 'selaku penerima kuasa dari {{ sponsor_company }}, dengan ini kami mengajukan '
           'permohonan', before=8)
    p(doc, '{{ body_request_line }}', after=2)

    field(doc, 'Nama', '{{ client_name }}')
    field(doc, 'Jenis Kelamin', '{{ client_gender }}')
    field(doc, 'Tempat/Tgl. Lahir', '{{ place_birth_inline }}')
    field(doc, 'Kewarganegaraan', '{{ client_nationality }}')
    field(doc, 'No. Paspor', '{{ passport_no }}')
    field(doc, 'Masa Berlaku Paspor', '{{ passport_expiry }}')
    field(doc, 'Nomor ITK', '{{ itk_no }}')
    field(doc, 'Masa Berlaku ITK', '{{ itk_expiry }}')
    field(doc, 'Alamat', '{{ client_address }}')

    p(doc, 'Demikian permohonan yang dapat kami ajukan, atas perhatiannya kami ucapkan '
           'terima kasih.', before=10, after=14)
    p(doc, '{{ place_date }}', align=WD_ALIGN_PARAGRAPH.RIGHT)
    p(doc, 'Hormat kami,', align=WD_ALIGN_PARAGRAPH.RIGHT, after=30)
    p(doc, '{{ attorney_name }}', align=WD_ALIGN_PARAGRAPH.RIGHT)

    tpl = ROOT / 'data' / 'templates' / code / 'word' / 'template.docx'
    tpl.parent.mkdir(parents=True, exist_ok=True)
    if tpl.exists():
        shutil.copy2(tpl, tpl.with_name(f'template.docx.bak_{datetime.now():%Y%m%d_%H%M}'))
    doc.save(tpl)
    print('собрано:', code, tpl.stat().st_size, 'байт')
    return tpl


def render(code: str) -> None:
    import sqlite3
    sys.path.insert(0, str(ROOT))
    from services.values import build_template_values
    from services import doc_engine

    c = sqlite3.connect(ROOT / 'data' / 'wolf.db')
    c.row_factory = sqlite3.Row
    r = c.execute('''SELECT a.id FROM applications a JOIN product_documents pd ON pd.product_id=a.product_id
                     JOIN templates t ON t.id=pd.template_id WHERE t.code=?
                     ORDER BY a.id DESC LIMIT 1''', (code,)).fetchone()
    app = r['id']
    a = dict(c.execute('SELECT * FROM applications WHERE id=?', (app,)).fetchone())
    cl = dict(c.execute('SELECT * FROM clients WHERE id=?', (a['client_id'],)).fetchone()) if a.get('client_id') else {}
    vals = build_template_values(cl, code, {}, app) or {}
    vals['client_address'] = re.sub(r'^[\s:;,\-]+', '', str(vals.get('client_address') or ''))
    res = doc_engine.build_document(code, vals, ROOT / 'data' / 'renders', engine='word',
                                    name=code.split('_')[0] + '_check')
    doc = pymupdf.open(res['pdf'])
    png = ROOT / 'data' / 'renders' / (code.split('_')[0] + '_check.png')
    doc[0].get_pixmap(dpi=110).save(png)
    print(f"  заявка {app}: страниц={len(doc)} пусто={res['missing']} ph={res['placeholders']} png={png.name}")


if __name__ == '__main__':
    for code in CODES:
        build(code)
    if '--render' in sys.argv:
        for code in CODES:
            render(code)
