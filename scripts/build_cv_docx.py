#!/usr/bin/env python3
"""Сборка Word-шаблона резюме (doc_index 5) по образцу resume_SABALEUSKAYA.docx.

В прежнем шаблоне было 4 переменные и 9 знаков статичного текста — фактически
пустой каркас. Здесь полная форма: имя, «Resume», должность, личные данные,
компетенции, опыт (2 места с обязанностями), образование, языки.
Имена переменных — те, что уже заведены в template_variables и отдаются
services/cv_content.py (cv_*).
"""
from __future__ import annotations

import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path

import docxtpl
import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ROOT = Path('/opt/wolf_v2')
CODE = 'cv__format_sivaev'
TPL = ROOT / 'data' / 'templates' / CODE / 'word' / 'template.docx'
GREY = RGBColor(0x44, 0x44, 0x44)


def p(doc, text='', size=10.5, bold=False, italic=False, caps=False, align=None,
      before=0, after=0, color=None, spacing=1.05):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), spacing
    if align is not None:
        par.alignment = align
    if text:
        r = par.add_run(text)
        r.font.all_caps = bool(caps)  # регистр — свойством шрифта, иначе имя переменной станет ЗАГЛАВНЫМ
        r.font.size, r.bold, r.italic = Pt(size), bold, italic
        r.font.name = 'Times New Roman'
        if color is not None:
            r.font.color.rgb = color
    return par


def rich(doc, value, size=10.5, after=0):
    """Абзац со значением-переменной; переносы строк сохраняются."""
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(0)
    par.paragraph_format.space_after = Pt(after)
    par.paragraph_format.line_spacing = 1.05
    r = par.add_run(value)
    r.font.size = Pt(size)
    r.font.name = 'Times New Roman'
    return par


def section(doc, title):
    p(doc, title, size=11, bold=True, caps=True, before=10, after=3)


def build() -> Path:
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(1.6)
    s.left_margin = s.right_margin = Cm(2.0)
    doc.styles['Normal'].font.name = 'Times New Roman'
    doc.styles['Normal'].font.size = Pt(10.5)

    p(doc, '{{ cv_name }}', size=19, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    p(doc, 'Resume', size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY, after=4)
    p(doc, '{{ cv_job1_title }}', size=11, bold=True, caps=True,
      align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

    # Личные данные
    for label, var in (('Date of Birth', 'cv_date_of_birth'),
                       ('Place of Birth', 'cv_place_of_birth'),
                       ('Nationality', 'cv_nationality'),
                       ('Passport No.', 'cv_passport_no')):
        par = doc.add_paragraph()
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1.05
        r1 = par.add_run(f'{label}: ')
        r1.bold = True
        r1.font.size = Pt(10.5)
        r1.font.name = 'Times New Roman'
        r2 = par.add_run('{{ %s }}' % var)
        r2.font.size = Pt(10.5)
        r2.font.name = 'Times New Roman'

    section(doc, 'Core competencies')
    rich(doc, '{{ cv_summary }}', after=3)
    rich(doc, '{{ cv_skills_hard }}', after=2)
    rich(doc, '{{ cv_skills_soft }}', after=2)
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(0)
    par.paragraph_format.space_after = Pt(0)
    r1 = par.add_run('Languages: ')
    r1.bold = True
    r1.font.size = Pt(10.5)
    r1.font.name = 'Times New Roman'
    r2 = par.add_run('{{ cv_languages }}')
    r2.font.size = Pt(10.5)
    r2.font.name = 'Times New Roman'

    section(doc, 'Work experience')
    p(doc, '{{ cv_job1_title }}', size=10.5, bold=True, after=0)
    p(doc, '{{ cv_job1_company }}{% if cv_job1_period %} | {{ cv_job1_period }}{% endif %}',
      size=10, italic=True, color=GREY, after=2)
    rich(doc, '{{ cv_job1_bullets }}', after=4)

    p(doc, '{{ cv_job2_title }}', size=10.5, bold=True, after=0)
    p(doc, '{{ cv_job2_company }}{% if cv_job2_period %} | {{ cv_job2_period }}{% endif %}',
      size=10, italic=True, color=GREY, after=2)
    rich(doc, '{{ cv_job2_bullets }}', after=4)

    section(doc, 'Education')
    p(doc, '{{ cv_education_degree }}', size=10.5, bold=True, after=0)
    p(doc, '{{ cv_education_school }}{% if cv_education_years %} | {{ cv_education_years }}{% endif %}',
      size=10, italic=True, color=GREY, after=0)

    TPL.parent.mkdir(parents=True, exist_ok=True)
    if TPL.exists():
        bak = TPL.with_name(f'template.docx.bak_{datetime.now():%Y%m%d_%H%M}')
        shutil.copy2(TPL, bak)
        print('бэкап:', bak.name)
    doc.save(TPL)
    print('собрано:', TPL, TPL.stat().st_size, 'байт')
    return TPL


def render(app: int = 109) -> None:
    import sqlite3
    sys.path.insert(0, str(ROOT))
    from services.values import build_template_values
    from services import doc_engine

    c = sqlite3.connect(ROOT / 'data' / 'wolf.db')
    c.row_factory = sqlite3.Row
    a = dict(c.execute('SELECT * FROM applications WHERE id=?', (app,)).fetchone())
    cl = dict(c.execute('SELECT * FROM clients WHERE id=?', (a['client_id'],)).fetchone())
    vals = build_template_values(cl, CODE, {}, app) or {}
    res = doc_engine.build_document(CODE, vals, ROOT / 'data' / 'renders',
                                    engine='word', name='cv_check')
    print('сборка:', {k: res[k] for k in ('ok', 'pages' if 'pages' in res else 'engine')}, '| пусто:', res['missing'],
          '| {{ }}:', res['placeholders'])
    pdf = pymupdf.open(res['pdf'])
    print('страниц:', len(pdf))
    pdf[0].get_pixmap(dpi=110).save(ROOT / 'data' / 'renders' / 'cv_check.png')
    print('png:', ROOT / 'data' / 'renders' / 'cv_check.png')


if __name__ == '__main__':
    build()
    if '--render' in sys.argv:
        render(107 if '107' in sys.argv else 109)
