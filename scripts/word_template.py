"""Единый конвейер: PDF-эталон → Word-шаблон (DOCX) с сохранением позиций.

Каждый текстовый фрагмент эталона кладётся в текстовый фрейм (w:framePr) с
абсолютными координатами от края страницы. Внутри фрейма текст обычный: длинное
значение переносится на следующую строку, шрифт/кегль/цвет сохраняются.
Картинки (логотип, печать) — плавающие якоря от страницы. Линии — тонкие якоря.

    word_template.py <code> [...]      один или несколько шаблонов
    word_template.py --all             все шаблоны (status не hidden/merged)
    word_template.py <code> --check    только сверка, без записи
"""
import os
import re
import shutil
import sqlite3
import subprocess
import sys

import docx
import pymupdf
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = '/opt/wolf_v2'
TEMPLATES = os.path.join(ROOT, 'data', 'templates')
VENDOR = os.path.join(ROOT, '.venv_docs', 'bin', 'python')
PT2EMU, TWIP = 12700, 20
Y_SHIFT = 0.0
ALT = {}                  # точечные поправки по строкам (нормализованный текст → pt)
A4_W, A4_H = 595.276, 841.89

FONT_MAP = (('Arial-BoldMT', 'Arial', True), ('ArialMT', 'Arial', False),
            ('Arial-BoldItalicMT', 'Arial', True), ('Arial-ItalicMT', 'Arial', False),
            ('Arimo', 'Arial', None), ('TimesNewRomanPS-BoldMT', 'Times New Roman', True),
            ('TimesNewRomanPSMT', 'Times New Roman', False), ('Helvetica', 'Arial', None),
            ('Courier', 'Courier New', None), ('Calibri', 'Calibri', None))


def font_of(span):
    name = span['font'].split('+')[-1] or 'Arial'
    family, bold = 'Arial', None
    for pat, fam, b in FONT_MAP:
        if pat in name:
            family, bold = fam, b
            break
    if bold is None:
        bold = 'Bold' in name
    return family, bold, 'Italic' in name


def color_of(span):
    c = span.get('color', 0)
    if isinstance(c, int) and c > 0:
        return RGBColor((c >> 16) & 255, (c >> 8) & 255, c & 255)
    return RGBColor(0, 0, 0)


def set_frame(par, x_pt, y_pt, w_pt, h_pt):
    """Абзац в абсолютном фрейме; x, y — от левого верхнего угла СТРАНИЦЫ."""
    pPr = par._p.get_or_add_pPr()
    frame = OxmlElement('w:framePr')
    frame.set(qn('w:w'), str(int(w_pt * TWIP)))
    frame.set(qn('w:h'), str(int(h_pt * TWIP)))
    frame.set(qn('w:hRule'), 'atLeast')
    frame.set(qn('w:wrap'), 'none')
    frame.set(qn('w:vAnchor'), 'page')
    frame.set(qn('w:hAnchor'), 'page')
    frame.set(qn('w:x'), str(int(x_pt * TWIP)))
    frame.set(qn('w:y'), str(int(y_pt * TWIP)))
    pPr.append(frame)
    par.paragraph_format.space_before = Pt(0)
    par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.line_spacing = Pt(1)


def anchor_picture(par, path, box):
    x, y, w, h = box
    run = par.add_run()
    run.add_picture(path, width=Pt(w), height=Pt(h))
    drawing = run._r.find(qn('w:drawing'))
    inline = drawing.find(qn('wp:inline'))
    anchor = OxmlElement('wp:anchor')
    for k, v in (('behindDoc', '0'), ('distT', '0'), ('distB', '0'), ('distL', '0'),
                 ('distR', '0'), ('simplePos', '0'), ('locked', '0'),
                 ('layoutInCell', '1'), ('allowOverlap', '1'), ('relativeHeight', '4')):
        anchor.set(k, v)
    sp = OxmlElement('wp:simplePos'); sp.set('x', '0'); sp.set('y', '0'); anchor.append(sp)
    for tag, val in (('wp:positionH', x), ('wp:positionV', y)):
        el = OxmlElement(tag); el.set('relativeFrom', 'page')
        off = OxmlElement('wp:posOffset'); off.text = str(int(val * PT2EMU))
        el.append(off); anchor.append(el)
    for ch in list(inline):
        if ch.tag == qn('wp:docPr'):
            anchor.append(OxmlElement('wp:wrapNone'))
        anchor.append(ch)
    drawing.remove(inline); drawing.append(anchor)


def extract(pdf):
    pages = []
    doc = pymupdf.open(pdf)
    for pg in doc:
        lines, images, rules = [], [], []
        for b in pg.get_text('dict')['blocks']:
            for ln in b.get('lines', []):
                spans = [s for s in ln['spans'] if s['text'].strip()]
                if spans:
                    lines.append({'bbox': ln['bbox'], 'spans': spans,
                                  'text': ''.join(s['text'] for s in ln['spans'])})
        for im in pg.get_image_info():
            images.append(im['bbox'])
        for d in pg.get_drawings():
            for item in d['items']:
                if item[0] == 'l':
                    p1, p2 = item[1], item[2]
                    if abs(p1.y - p2.y) < 1.0 and abs(p2.x - p1.x) > 40:
                        rules.append((min(p1.x, p2.x), p1.y, max(p1.x, p2.x)))
        pages.append({'lines': lines, 'images': images, 'rules': rules})
    return pages


def build(code, src_pdf, out_docx, vars_map=None, footer_text=None):
    workdir = os.path.join('/tmp', 'wordbuild_' + code)
    os.makedirs(workdir, exist_ok=True)
    pages = extract(src_pdf)
    _mark_vals = [v for v in (vars_map if vars_map is not None else vars_from_db(code))
                  if len(v) >= 3]
    _chunks = markup_from_html(code) if vars_map is not None else []
    src = pymupdf.open(src_pdf)
    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin, sec.right_margin = Cm(1.83), Cm(1.67)
    sec.top_margin, sec.bottom_margin = Cm(0.5), Cm(0.5)
    sec.header_distance, sec.footer_distance = Cm(0.4), Cm(1.2)
    if footer_text:
        fr = sec.footer.paragraphs[0].add_run(footer_text)
        fr.font.size, fr.font.name = Pt(9.0), 'Arial'
        fr.font.color.rgb = RGBColor(0x90, 0x90, 0x90)

    for pi, page in enumerate(pages):
        if pi:
            doc.add_page_break()
        for i, (x0, y0, x1, y1) in enumerate(page['images']):
            if x1 - x0 < 2 or y1 - y0 < 2:
                continue
            png = os.path.join(workdir, 'img_%d_%d.png' % (pi, i))
            if not os.path.exists(png):
                src[pi].get_pixmap(dpi=300, clip=pymupdf.Rect(x0, y0, x1, y1)).save(png)
            anchor_picture(doc.add_paragraph(), png, (x0, y0, x1 - x0, y1 - y0))
        for j, (x0, y0, x1) in enumerate(page['rules']):
            png = os.path.join(workdir, 'rule_%d_%d.png' % (pi, j))
            if not os.path.exists(png):
                src[pi].get_pixmap(dpi=300,
                                   clip=pymupdf.Rect(x0, y0 - 0.6, x1, y0 + 0.6)).save(png)
            anchor_picture(doc.add_paragraph(), png, (x0, y0 - 0.6, x1 - x0, 1.2))
        # «стопки» одной колонки: строки-соседи идут одним потоком, чтобы длинное
        # значение переменной сдвигало их вниз, а не накладывалось на них
        stack_of, in_stack = {}, {}
        order = sorted(range(len(page['lines'])),
                       key=lambda i: (round(page['lines'][i]['bbox'][0], 1),
                                      page['lines'][i]['bbox'][1]))
        used = set()
        for i in order:
            if i in used:
                continue
            stack, used = [i], used | {i}
            base = page['lines'][i]
            for j in order:
                if j in used:
                    continue
                lj = page['lines'][j]
                ov = (min(lj['bbox'][2], base['bbox'][2])
                      - max(lj['bbox'][0], base['bbox'][0]))
                narrow = min(lj['bbox'][2] - lj['bbox'][0], base['bbox'][2] - base['bbox'][0])
                if ov < 0.3 * narrow:            # строки должны стоять друг над другом
                    continue
                prev = page['lines'][stack[-1]]
                gap = lj['bbox'][1] - prev['bbox'][3]
                if -8.0 <= gap <= 8.0:           # строки вплотную друг под другом
                    stack.append(j)
                    used.add(j)
            _ys2 = [page['lines'][k]['bbox'][1] for k in stack]
            _hs2 = [page['lines'][k]['bbox'][3] - page['lines'][k]['bbox'][1] for k in stack]
            _steps = [_ys2[i + 1] - _ys2[i] for i in range(len(_ys2) - 1)]
            _even = bool(_steps) and (max(_steps) - min(_steps)) <= 2.0
            _tight = all(s >= 0.85 * h for s, h in zip(_steps, _hs2))
            _x0s = [page['lines'][k]['bbox'][0] for k in stack]
            _same_x = (max(_x0s) - min(_x0s)) <= 3.0
            if (2 <= len(stack) <= 4 and _even and _tight and _same_x
                    and any(v in page['lines'][k]['text'] for k in stack for v in _mark_vals)):
                stack_of[stack[0]] = stack
                for k in stack[1:]:
                    in_stack[k] = stack[0]

        # Абзацы: соседние строки одного столбца, которые вместе совпадают с картой HTML.
        # Нужны затем, что в PDF абзац часто разбит на строки, а переменная стоит в середине
        # («... is made on this {{ agreement_date }}.») — по одной строке такое место не найти.
        if _chunks:
            _in_para = {}
            for _i in range(len(page['lines'])):
                if any(_i in v for v in _in_para.values()) or (page['lines'][_i].get('stack')):
                    continue
                _grp, _cur = [_i], _i
                while True:
                    _nxt = None
                    for _j in range(_cur + 1, len(page['lines'])):
                        if _j in _in_para or page['lines'][_j].get('stack'):
                            continue
                        _a, _b = page['lines'][_cur], page['lines'][_j]
                        if (abs(_a['bbox'][0] - _b['bbox'][0]) <= 3.0
                                and 0 <= _b['bbox'][1] - _a['bbox'][3] <= 8.0):
                            _nxt = _j
                            break
                    if _nxt is None:
                        break
                    _grp.append(_nxt)
                    _in_para[_nxt] = _grp
                    _cur = _nxt
                if len(_grp) < 2:
                    continue
                _probe = markup_segments(' '.join(page['lines'][k]['text'] for k in _grp), _chunks)
                # абзац годится, только если в разметке есть хотя бы одна переменная:
                # иначе склейка «съела» бы строку с переменной, заменив её статикой
                if not _probe or not any(_v for _t, _v, _p in _probe):
                    continue
                page['lines'][_i]['para'] = _grp
                for k in _grp[1:]:
                    page['lines'][k]['skip'] = True

        for li, ln in enumerate(page['lines']):
            if ln.get('skip'):
                continue
            if li in in_stack:
                continue
            if li in stack_of:
                stack = stack_of[li]
                ln = dict(page['lines'][li])
                ln['text'] = '\n'.join(page['lines'][k]['text'] for k in stack)
                ln['stack'] = stack
                ln['bbox'] = (ln['bbox'][0], ln['bbox'][1],
                              max(page['lines'][k]['bbox'][2] for k in stack),
                              ln['bbox'][3])
            x0, y0, x1, y1 = ln['bbox']
            h = max(10.0, y1 - y0) + 1.0
            w = min(max(30.0, (x1 - x0) + 60.0), A4_W - x0 - 6.0)
            par = doc.add_paragraph()
            key = norm(ln['text'])
            yfix = globals().get('ALT', {}).get(key, 0.0)
            set_frame(par, x0, max(0.0, y0 + globals().get('Y_SHIFT', 0.0) + yfix), w, h)
            _sz = max(sp['size'] for sp in ln['spans'])
            par.paragraph_format.line_spacing = Pt(round(_sz * 1.16, 1))
            if ln.get('stack'):
                _ys = [page['lines'][k]['bbox'][1] for k in ln['stack']]
                _step = (max(_ys) - min(_ys)) / max(1, len(_ys) - 1)
                par.paragraph_format.line_spacing = Pt(round(_step, 1))
                for _si, _k in enumerate(ln['stack']):
                    _prev = None
                    for _sp in page['lines'][_k]['spans']:
                        _t = _sp['text']
                        if '{{' not in _t:
                            for _s, _d in (vars_map or {}).items():
                                if _s and _s in _t:
                                    _t = _t.replace(_s, _d)
                        if _prev is not None:
                            _gp = _sp['bbox'][0] - _prev['bbox'][2]
                            if (_gp > 1.2 and not _prev['text'].endswith(' ')
                                    and not _t.startswith(' ')):
                                _t = ' ' + _t
                        _prev = _sp
                        _r = par.add_run(_t)
                        _fm, _bd, _it = font_of(_sp)
                        _r.font.size, _r.font.name = Pt(round(_sp['size'], 1)), _fm
                        _r.bold, _r.italic = _bd, _it
                        _r.font.color.rgb = color_of(_sp)
                    if _si + 1 < len(ln['stack']):
                        par.add_run().add_break()
                prev_bottom = page['lines'][ln['stack'][-1]]['bbox'][3]
                continue
            if ln.get('para'):
                _grp = ln['para']
                _ys = [page['lines'][k]['bbox'][1] for k in _grp]
                _steps = [_ys[i + 1] - _ys[i] for i in range(len(_ys) - 1)]
                _step = sum(_steps) / len(_steps) if _steps else 10.0
                par.paragraph_format.line_spacing = Pt(round(_step, 1))
                _segs_p = markup_segments(' '.join(page['lines'][k]['text'] for k in _grp), _chunks)
                _sp0 = ln['spans'][0] if ln['spans'] else None
                for _txt, _isvar, _pos in (_segs_p or []):
                    _r = par.add_run(_txt)
                    if _sp0:
                        _fm, _bd, _it = font_of(_sp0)
                        _r.font.size, _r.font.name = Pt(round(_sp0['size'], 1)), _fm
                        _r.bold, _r.italic = _bd, _it
                        _r.font.color.rgb = color_of(_sp0)
                prev_bottom = page['lines'][_grp[-1]]['bbox'][3]
                continue
            _segs = markup_segments(ln['text'], _chunks) if _chunks else None
            if _segs:
                _spans = ln['spans']
                _bounds, _cum = [], 0
                for _sp in _spans:
                    _bounds.append((_cum, _cum + len(_sp['text'])))
                    _cum += len(_sp['text'])

                def _sp_at(_p):
                    for _b, _s in zip(_bounds, _spans):
                        if _b[0] <= _p < _b[1]:
                            return _s
                    return _spans[0]

                for _txt, _isvar, _pos in _segs:
                    _sp = _sp_at(_pos)
                    _r = par.add_run(_txt)
                    _fm, _bd, _it = font_of(_sp)
                    _r.font.size, _r.font.name = Pt(round(_sp['size'], 1)), _fm
                    _r.bold, _r.italic = _bd, _it
                    _r.font.color.rgb = color_of(_sp)
                prev_bottom = y1
                continue
            prev_sp = None
            for sp in ln['spans']:
                txt = sp['text']
                if vars_map and '{{' not in txt:
                    for s, d in vars_map.items():
                        if s and s in txt:
                            txt = txt.replace(s, d)
                # восстановить пробел, если в PDF фрагменты разнесены
                if prev_sp is not None:
                    gap = sp['bbox'][0] - prev_sp['bbox'][2]
                    if gap > 1.2 and not prev_sp['text'].endswith(' ') and not txt.startswith(' '):
                        txt = ' ' + txt
                prev_sp = sp
                r = par.add_run(txt)
                fam, bold, italic = font_of(sp)
                r.font.size, r.font.name = Pt(round(sp['size'], 1)), fam
                r.bold, r.italic = bold, italic
                r.font.color.rgb = color_of(sp)
    doc.save(out_docx)
    return out_docx


def render_pdf(docx_path):
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', docx_path,
                    '--outdir', os.path.dirname(docx_path)], capture_output=True, timeout=300)
    return os.path.splitext(docx_path)[0] + '.pdf'


def norm(t):
    return re.sub(r'\s+', ' ', t).strip().lower()


def rows(pdf):
    out = []
    for b in pymupdf.open(pdf)[0].get_text('dict')['blocks']:
        for l in b.get('lines', []):
            t = ''.join(s['text'] for s in l['spans']).strip()
            if t:
                out.append({'x': round(l['bbox'][0], 1), 'y': round(l['bbox'][1], 1),
                            'text': t, 'key': norm(t)})
    return out


def compare(orig_pdf, our_pdf):
    """Сопоставление по тексту строки (не по порядку): сверка начала строки."""
    a, b = rows(orig_pdf), rows(our_pdf)
    used, bad, missing = set(), [], []
    for ra in a:
        hit = None
        for j, rb in enumerate(b):
            if j in used:
                continue
            if ra['key'] == rb['key'] or (len(ra['key']) > 6 and
                                          (ra['key'].startswith(rb['key']) or rb['key'].startswith(ra['key']))):
                hit = j
                break
        if hit is None:
            missing.append(ra['text'][:48])
            continue
        used.add(hit)
        rb = b[hit]
        dx, dy = ra['x'] - rb['x'], ra['y'] - rb['y']   # знаковые: нужны для подгонки
        if abs(dx) > 2 or abs(dy) > 2:
            bad.append((round(abs(dx)), round(dy), ra['key'], rb['key']))
    extra = [rb['text'][:48] for j, rb in enumerate(b) if j not in used]
    return {'orig': len(a), 'our': len(b), 'bad': bad, 'missing': missing, 'extra': extra}


def markup_from_html(code):
    """Строки HTML-шаблона с {{переменными}} — источник точной разметки."""
    import re as _re
    import html as _html
    h = os.path.join(TEMPLATES, code, 'template.html')
    if not os.path.exists(h):
        return []
    raw = open(h, encoding='utf-8', errors='ignore').read()
    raw = _re.sub(r'<(script|style)[^>]*>.*?</\1>', ' ', raw, flags=_re.S | _re.I)
    txt = _html.unescape(_re.sub(r'<[^>]+>', ' ', raw))
    out = []
    for chunk in _re.split(r'(?<=[.;:])\s|[\n\r]+', txt):
        chunk = _re.sub(r'\s+', ' ', chunk).strip()
        if '{{' in chunk and len(chunk) > 6:
            out.append(chunk)
    return out


def apply_markup(line_text, chunks):
    """Заменяет значения в строке по разметке HTML (regex по статичным частям)."""
    import re as _re
    flat = _re.sub(r'\s+', ' ', line_text).strip()
    for ch in chunks:
        parts = _re.split(r'\{\{\s*([a-zA-Z0-9_]+)\s*\}\}', ch)
        rx, names, out = '', [], ''
        ok = True
        for i, part in enumerate(parts):
            if i % 2 == 0:
                st = _re.sub(r'\s+', ' ', part)
                if st.strip(' ,.;:') == '' and i not in (0, len(parts) - 1):
                    st = ''
                rx += _re.escape(st)
                out += st
            else:
                rx += '(.+?)'
                names.append(part)
        try:
            m = _re.fullmatch(rx, flat)
        except Exception:
            m = None
        if m:
            res, g = '', 0
            for i, part in enumerate(parts):
                if i % 2 == 0:
                    res += _re.sub(r'\s+', ' ', part)
                else:
                    res += '{{%s}}' % names[g]
                    g += 1
            return res.strip()
    return line_text


def markup_segments(line_text, chunks):
    """Разметка строки по HTML-шаблону: [(текст, признак_переменной, позиция_в_строке)]."""
    import re as _re
    flat = _re.sub(r'\s+', ' ', line_text).strip()
    for ch in chunks:
        parts = _re.split(r'\{\{\s*([a-zA-Z0-9_]+)\s*\}\}', ch)
        rx = ''
        for i, part in enumerate(parts):
            rx += _re.escape(_re.sub(r'\s+', ' ', part)) if i % 2 == 0 else '(.+?)'
        try:
            m = _re.fullmatch(rx, flat)
        except Exception:
            m = None
        if not m:
            continue
        segs, pos = [], 0
        for i, part in enumerate(parts):
            if i % 2 == 0:
                st = _re.sub(r'\s+', ' ', part)
                segs.append((st, False, pos))
                pos += len(st)
            else:
                val = m.groups()[i // 2]
                segs.append(('{{%s}}' % part, True, pos))
                pos += len(val)
        return segs
    # второй проход: без пунктуации (в карте «G. Sochi», в документе «G Sochi»)
    _lax = lambda s: _re.sub(r'\s+', ' ', _re.sub(r'[^0-9A-Za-z\u0410-\u044f]+', ' ', s)).strip().lower()
    flat_l = _lax(line_text)
    for ch in chunks:
        parts = _re.split(r'\{\{\s*([a-zA-Z0-9_]+)\s*\}\}', ch)
        rx_l = ''
        for i, part in enumerate(parts):
            rx_l += _re.escape(_lax(part)) if i % 2 == 0 else '(.+?)'
        try:
            m2 = _re.fullmatch(rx_l, flat_l)
        except Exception:
            m2 = None
        if not m2:
            continue
        segs, pos = [], 0
        for i, part in enumerate(parts):
            if i % 2 == 0:
                st = _re.sub(r'\s+', ' ', part)
                segs.append((st, False, pos))
                pos += len(st)
            else:
                segs.append(('{{%s}}' % part, True, pos))
                pos += len(m2.groups()[i // 2])
        return segs
    return None


def vars_from_db(code):
    """Значения переменных из базы → карта замен «значение в эталоне» → «{{имя}}»."""
    import json
    con = sqlite3.connect(os.path.join(ROOT, 'data', 'wolf.db'))
    row = con.execute("SELECT id FROM templates WHERE code=?", (code,)).fetchone()
    if not row:
        return {}
    known = set()
    for r in con.execute("SELECT name FROM template_variables WHERE template_id=?", (row[0],)):
        known.add(r[0])
    m = {}
    for name, dv, ex, op in con.execute(
            "SELECT name, default_value, examples_json, options_json FROM template_variables"
            " WHERE template_id=?", (row[0],)):
        vals = []
        if dv:
            vals.append(str(dv))
        for js in (ex or '', op or ''):
            try:
                vals += [str(v) for v in json.loads(js)]
            except Exception:
                pass
        for v in vals:
            v = v.strip()
            if len(v) >= 3 and v not in known:      # не заменять имена переменных
                m[v] = '{{%s}}' % name
    # сначала самые длинные значения, чтобы не портить вложенные совпадения
    return dict(sorted(m.items(), key=lambda kv: -len(kv[0])))


def process(code, save=True):
    d = os.path.join(TEMPLATES, code)
    src = os.path.join(d, 'example_filled.pdf')
    if not os.path.exists(src):
        print('  %-40s НЕТ эталона' % code)
        return None
    outdir = os.path.join(d, 'word')
    os.makedirs(outdir, exist_ok=True)
    docx_path = os.path.join(outdir if save else '/tmp', 'template.docx' if save else 'chk_%s.docx' % code)
    global Y_SHIFT
    Y_SHIFT = 0.0                     # важно: поправка не переносится на следующий документ
    build(code, src, docx_path)
    pdf = render_pdf(docx_path)
    # автокалибровка вертикали: пробуем оба направления, оставляем лучший результат
    best_shift = 0.0
    best_bad = len(compare(src, pdf)['bad'])
    for _ in range(3):
        r0 = compare(src, pdf)
        dys = sorted(w[1] for w in r0['bad'] if w[0] == 0)
        if len(dys) < 3:
            break
        med = dys[len(dys) // 2]
        if abs(med) <= 1.5:
            break
        improved = False
        for cand in (best_shift + med, best_shift - med):
            Y_SHIFT = cand
            build(code, src, docx_path)
            pdf = render_pdf(docx_path)
            r = compare(src, pdf)
            if len(r['bad']) < best_bad:
                best_shift, best_bad, improved = cand, len(r['bad']), True
                break
        if not improved:
            break
    Y_SHIFT = best_shift
    ALT.clear()
    # точечная доводка: строки, которые всё ещё смещены, правим поимённо
    for _ in range(3):
        build(code, src, docx_path)
        pdf = render_pdf(docx_path)
        r = compare(src, pdf)
        if not r['bad']:
            break
        a = {x['key']: x['y'] for x in rows(src)}
        b = {x['key']: x['y'] for x in rows(pdf)}
        changed = False
        for _dx, _dy, key_o, key_u in r['bad']:
            k = key_o or key_u
            if k in a and k in b and abs(a[k] - b[k]) > 1.0:
                ALT[k] = ALT.get(k, 0.0) + (a[k] - b[k])
                changed = True
        if not changed:
            break
    res = compare(src, pdf)
    cal_pdf = pdf                                  # версия «как эталон» — для сверки наложением
    vm = vars_from_db(code)                         # подключение переменных шаблона
    build(code, src, docx_path, vars_map=vm if vm else {})
    pdf = render_pdf(docx_path)
    _nvar = 0
    try:
        import zipfile as _zf
        _doc = _zf.ZipFile(docx_path).read('word/document.xml').decode()
        _nvar = len(set(re.findall(r'\{\{\s*([a-zA-Z0-9_]+)\s*\}\}', _doc)))
    except Exception:
        pass
    print('        подключено переменных: %d' % _nvar)
    ov = '/tmp/word_%s_overlay.png' % code
    subprocess.run([VENDOR, os.path.join(ROOT, 'scripts', 'overlay_diff.py'),
                    src, cal_pdf, ov, '100'], capture_output=True, timeout=240)
    pages = pymupdf.open(pdf).page_count
    print('  %-40s строк %d/%d, страниц %d, сдвинуто %d, не найдено %d, лишних %d'
          % (code, res['our'], res['orig'], pages, len(res['bad']), len(res['missing']), len(res['extra'])))
    for w in sorted(res['bad'], key=lambda w: -(w[0] + w[1]))[:6]:
        print('        Δx=%s Δy=%s | %r | %r' % w)
    for m in res['missing'][:4]:
        print('        НЕТ: %r' % m)
    for e in res['extra'][:4]:
        print('        ЛИШНЕЕ: %r' % e)
    if save:
        shutil.copy(pdf, os.path.join(outdir, 'preview.pdf'))
    return res


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    flags = {a for a in sys.argv[1:] if a.startswith('--')}
    if '--all' in flags:
        con = sqlite3.connect(os.path.join(ROOT, 'data', 'wolf.db'))
        codes = [r[0] for r in con.execute(
            "SELECT code FROM templates WHERE COALESCE(status,'ok') NOT IN ('hidden','merged')"
            " ORDER BY COALESCE(doc_index,0), id")]
    else:
        codes = args
    if not codes:
        print(__doc__)
        return
    print('=== PDF эталон → Word-шаблон ===')
    for c in codes:
        process(c, save='--check' not in flags)
    print('готово')


if __name__ == '__main__':
    main()
