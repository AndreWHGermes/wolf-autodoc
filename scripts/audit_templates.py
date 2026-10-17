#!/usr/bin/env python3
"""Сплошная ревизия шаблонов Wolf: чем каждый собирается и чего в нём не хватает.

Для каждого активного шаблона берём заявку его продукта, считаем значения тем же
путём, что у сайта, собираем документ и смотрим:
  * сколько статичного текста в шаблоне (пустой каркас = только переменные);
  * сколько переменных и какие остались пустыми;
  * остались ли в PDF незакрытые {{ }};
  * сколько страниц вышло.

Отчёт: /root/.hermes/cache/wolf_audit/report.md и report.json
"""
from __future__ import annotations

import json
import re
import sqlite3
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, '/opt/wolf_v2')
import pymupdf
from docx import Document

from services.values import build_template_values
from services import doc_engine

OUT = Path('/root/.hermes/cache/wolf_audit')
OUT.mkdir(parents=True, exist_ok=True)
DB = '/opt/wolf_v2/data/wolf.db'
TPL_ROOT = Path('/opt/wolf_v2/data/templates')

c = sqlite3.connect(DB)
c.row_factory = sqlite3.Row

tpls = [dict(r) for r in c.execute(
    'SELECT id, code, title, doc_index, engine, status FROM templates ORDER BY doc_index, id')]
tpls = [t for t in tpls if (t['status'] or 'ok') in ('ok', 'active', 'published') and t['engine']]

rows = []
for t in tpls:
    code = t['code']
    d = TPL_ROOT / code
    item = {
        'doc_index': t['doc_index'], 'code': code, 'title': t['title'],
        'engine': t['engine'], 'vars': 0, 'static_chars': 0,
        'pages': None, 'missing': [], 'placeholders': [], 'app': None, 'error': None,
    }
    # каркас
    if t['engine'] == 'word':
        wp = d / 'word' / 'template.docx'
        if wp.exists():
            txt = '\n'.join(p.text for p in Document(str(wp)).paragraphs)
            names = set(re.findall(r'\{\{\s*([A-Za-z_]\w*)\s*\}\}', txt))
            item['vars'] = len(names)
            item['static_chars'] = len(re.sub(r'\{\{[^}]*\}\}', '', txt).strip())
    else:
        hp = d / 'template.html'
        if hp.exists():
            raw = hp.read_text(encoding='utf-8', errors='ignore')
            body = re.sub(r'<[^>]+>', ' ', raw)
            names = set(re.findall(r'\{\{\s*([A-Za-z_]\w*)\s*\}\}', body))
            item['vars'] = len(names)
            item['static_chars'] = len(re.sub(r'\{\{[^}]*\}\}', '', body).strip())

    # заявка продукта
    app = c.execute(
        """SELECT a.id FROM applications a
           JOIN product_documents pd ON pd.product_id = a.product_id
           JOIN templates t2 ON t2.id = pd.template_id
           WHERE t2.id = ? ORDER BY a.id DESC LIMIT 1""", (t['id'],)).fetchone()
    if app:
        app_id = app['id']
        item['app'] = app_id
        a = dict(c.execute('SELECT * FROM applications WHERE id=?', (app_id,)).fetchone())
        cl = {}
        if a.get('client_id'):
            cl = dict(c.execute('SELECT * FROM clients WHERE id=?', (a['client_id'],)).fetchone())
        try:
            vals = build_template_values(cl, code, {}, app_id) or {}
            res = doc_engine.build_document(code, vals, OUT / 'build', engine=t['engine'],
                                            name=f"{t['doc_index']:02d}_{code}")
            item['missing'] = [v for v in (res.get('missing') or [])]
            item['placeholders'] = list(res.get('placeholders') or [])[:6]
            if res.get('pdf'):
                doc = pymupdf.open(res['pdf'])
                item['pages'] = len(doc)
                item['pdf_chars'] = sum(len(p.get_text()) for p in doc)
        except Exception as e:  # noqa: BLE001
            item['error'] = f'{type(e).__name__}: {e}'[:160]
    rows.append(item)

rows.sort(key=lambda r: (r['doc_index'] or 99, r['code']))
(OUT / 'report.json').write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding='utf-8')

lines = ['# Ревизия шаблонов Wolf', '',
         '| № | шаблон | движок | перем. | статич. знаков | стр. | пусто | {{ }} | заявка |',
         '|---|---|---|---|---|---|---|---|---|']
for r in rows:
    lines.append('| {d} | {c} | {e} | {v} | {s} | {p} | {m} | {ph} | {a} |'.format(
        d=r['doc_index'], c=r['code'], e=r['engine'], v=r['vars'], s=r['static_chars'],
        p=r['pages'] if r['pages'] is not None else '—',
        m=len(r['missing']), ph=len(r['placeholders']), a=r['app'] or '—'))
(OUT / 'report.md').write_text('\n'.join(lines), encoding='utf-8')

for r in rows:
    print(f"{r['doc_index']:>2} {r['code'][:42]:<42} {r['engine']:<5} перем={r['vars']:<3} "
          f"стат={r['static_chars']:<5} стр={r['pages']} пусто={len(r['missing'])} "
          f"ph={len(r['placeholders'])} {('ОШИБКА ' + r['error']) if r['error'] else ''}")
print('\nотчёт:', OUT / 'report.md')
