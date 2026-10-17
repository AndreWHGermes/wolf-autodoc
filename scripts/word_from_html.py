#!/usr/bin/env python3
"""HTML-шаблон (один слой текста) → Word-шаблон DOCX + прогон значений.

Стандарт: каждый абзац/строка эталона — ОДИН текстовый фрейм, внутри которого
статика и {{переменные}} вместе. Word сам переносит текст и сдвигает строки
внутри фрейма, поэтому значение любой длины не налезает на соседний блок —
это то, чего не давали прежние шаблоны (статика и значения лежали в разных
фреймах, и длинное значение перекрывало чужой текст).

Подложка (логотип, линии, печать, подпись) — картинка под текстом (behindDoc).

    word_from_html.py <code> [--app 89] [--out-docx]
"""
import argparse
import html as H
import importlib.util
import os
import re
import subprocess
import sys

ROOT = "/opt/wolf_v2"
TWIP = 20
PT2EMU = 12700
RIGHT = 549.0
FONT = "Liberation Sans"
OUT = "/tmp/wordrun"

sys.path.insert(0, ROOT)
from docx import Document                      # noqa: E402
from docx.shared import Pt, RGBColor           # noqa: E402
from docx.oxml import OxmlElement              # noqa: E402
from docx.oxml.ns import qn                    # noqa: E402
from docxtpl import DocxTemplate               # noqa: E402


def set_frame(par, x, y, w, h):
    """Абзац в абсолютном фрейме от угла СТРАНИЦЫ; высота растёт (hRule=atLeast)."""
    pPr = par._p.get_or_add_pPr()
    fr = OxmlElement("w:framePr")
    fr.set(qn("w:w"), str(int(w * TWIP)))
    fr.set(qn("w:h"), str(int(h * TWIP)))
    fr.set(qn("w:hRule"), "atLeast")
    fr.set(qn("w:wrap"), "none")
    fr.set(qn("w:vAnchor"), "page")
    fr.set(qn("w:hAnchor"), "page")
    fr.set(qn("w:x"), str(int(x * TWIP)))
    fr.set(qn("w:y"), str(int(y * TWIP)))
    pPr.append(fr)
    par.paragraph_format.space_before = Pt(0)
    par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.line_spacing = 1.3


def anchor_picture(par, path, box, behind=True):
    x, y, w, h = box
    run = par.add_run()
    run.add_picture(path, width=Pt(w), height=Pt(h))
    drawing = run._r.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))
    anchor = OxmlElement("wp:anchor")
    for k, v in (("behindDoc", "1" if behind else "0"), ("distT", "0"), ("distB", "0"),
                 ("distL", "0"), ("distR", "0"), ("simplePos", "0"), ("locked", "0"),
                 ("layoutInCell", "1"), ("allowOverlap", "1"), ("relativeHeight", "1")):
        anchor.set(k, v)
    sp = OxmlElement("wp:simplePos"); sp.set("x", "0"); sp.set("y", "0"); anchor.append(sp)
    for tag, val in (("wp:positionH", x), ("wp:positionV", y)):
        el = OxmlElement(tag); el.set("relativeFrom", "page")
        off = OxmlElement("wp:posOffset"); off.text = str(int(val * PT2EMU))
        el.append(off); anchor.append(el)
    for ch in list(inline):
        if ch.tag == qn("wp:docPr"):
            anchor.append(OxmlElement("wp:wrapNone"))
        anchor.append(ch)
    drawing.remove(inline); drawing.append(anchor)


def read_blocks(code):
    """Абзацы и точки из HTML-шаблона: координаты, кегль, жирность, текст с {{переменными}}."""
    path = os.path.join(ROOT, "data", "templates", code, "template.html")
    src = open(path, encoding="utf-8").read()
    out = []
    for style, inner in re.findall(r'<span class="v" style="([^"]*)"[^>]*>(.*?)</span>', src, re.S):
        st = dict(re.findall(r"([a-z\-]+)\s*:\s*([0-9.]+)pt", style))
        col = re.search(r"color:\s*(#[0-9a-fA-F]{6})", style)
        x = float(st.get("left", 40.0))
        y = float(st.get("top", 40.0))
        size = float(st.get("font-size", 9.0))
        w = float(st.get("width", RIGHT - x)) or (RIGHT - x)
        txt = H.unescape(re.sub(r"<[^>]+>", "", inner)).strip()
        if txt:
            out.append(dict(x=x, y=y, w=min(w, RIGHT - x), size=size,
                            bold=("font-weight:700" in style), text=txt,
                            color=(col.group(1) if col else "#000000")))
    return out


def build(code, out_docx):
    blocks = read_blocks(code)
    doc = Document()
    sec = doc.sections[0]
    # размер листа — из оригинала (E-02 печатается на 612.5×936.5 pt, не на A4): иначе
    # подложка растягивается и подписи бланка уезжают относительно значений
    import pymupdf as _fitz
    pw, ph = 595.28, 841.89
    for _cand in ("blank.pdf", "example_filled.pdf"):
        _p = os.path.join(ROOT, "data", "templates", code, _cand)
        if os.path.exists(_p):
            _d = _fitz.open(_p)
            pw, ph = float(_d[0].rect.width), float(_d[0].rect.height)
            _d.close()
            break
    for attr, val in (("page_width", pw), ("page_height", ph),
                      ("left_margin", 0), ("right_margin", 0), ("top_margin", 0), ("bottom_margin", 0)):
        setattr(sec, attr, Pt(val))

    # подложка: вырезанный бланк (полосы абзацев уже стёрты при сборке HTML)
    bg = None
    for name in ("bg_para.png", "bg.png"):
        p = os.path.join(ROOT, "data", "templates", code, "frame", name)
        if os.path.exists(p):
            bg = p
            break
    # Для резюме фон-образец не подкладываем: на bg.png нарисован заполненный
    # чужой документ (SIVAEV, Executive Director, INGREDIENTTRADE), и наш текст
    # печатался поверх него — в визоре выходила «каша» из двух документов
    # (Андрей, 17.10.2026). Резюме — текстовый документ без логотипа и печати.
    if bg and not str(code).startswith("cv"):
        anchor_picture(doc.add_paragraph(), bg, (0.0, 0.0, pw, ph))

    for b in blocks:
        par = doc.add_paragraph()
        set_frame(par, b["x"], b["y"], b["w"], b["size"] * 1.45)
        run = par.add_run(b["text"])
        run.font.name = FONT
        run.font.size = Pt(b["size"])
        run.font.bold = b["bold"]
        run.font.color.rgb = RGBColor.from_string((b.get("color") or "#000000").lstrip("#").upper())
        rpr = run._r.get_or_add_rPr()
        rf = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            rf.set(qn(a), FONT)
        rpr.append(rf)
    doc.save(out_docx)
    return len(blocks)


def values_for(app_id, code):
    spec = importlib.util.spec_from_file_location("pdf_fill_mod", os.path.join(ROOT, "scripts", "pdf_fill.py"))
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m.values_for(app_id, code) or {}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("code")
    ap.add_argument("--app", type=int)
    ap.add_argument("--out-docx", action="store_true", help="положить рядом с шаблоном (word/template_flow.docx)")
    a = ap.parse_args()

    if str(a.code).startswith("cv__format"):
        # Резюме — потоковый документ: Word-версию собираем из того же HTML, что и просмотр
        # (scripts/cv_flow.py). Координатные фреймы с длинными описаниями разъезжаются.
        cmd = [sys.executable, os.path.join(os.path.dirname(os.path.abspath(__file__)), "cv_flow.py"),
               a.code, "--docx"]
        if a.app:
            cmd += ["--app", str(a.app)]
        return subprocess.run(cmd, check=False).returncode

    os.makedirs(OUT, exist_ok=True)
    tpl_docx = os.path.join(OUT, "%s_tpl.docx" % a.code)
    n = build(a.code, tpl_docx)
    print("блоков в шаблоне: %d → %s" % (n, tpl_docx))
    if a.out_docx:
        dst = os.path.join(ROOT, "data", "templates", a.code, "word", "template_flow.docx")
        os.replace(tpl_docx, dst)
        print("шаблон сохранён:", dst)

    if a.app:
        filled = os.path.join(OUT, "%s_app%s.docx" % (a.code, a.app))
        t = DocxTemplate(tpl_docx)
        vals = {k: ("" if v is None else str(v)) for k, v in values_for(a.app, a.code).items()}
        # В документ идёт только английский: русская подпись профессии нужна в интерфейсе,
        # в бланке она лишняя («Cinematographer, Film Editor (Киноператор и киномонтажер)»).
        import re as _re
        for _k, _val in list(vals.items()):
            if _k.startswith("cv_") and _re.search(r"\([^)]*[А-Яа-яЁё][^)]*\)", _val):
                vals[_k] = _re.sub(r"\s*\([^)]*[А-Яа-яЁё][^)]*\)", "", _val)
        t.render(vals)
        t.save(filled)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", OUT, filled],
                       capture_output=True, timeout=240)
        print("прогон:", filled.replace(".docx", ".pdf"))


if __name__ == "__main__":
    main()
