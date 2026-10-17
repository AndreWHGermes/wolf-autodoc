#!/usr/bin/env python3
"""Сборка Word-шаблона письма Bridging (doc_index 9) по подлиннику.

Проблема, которую решает: прежний word/template.docx содержал только строки
переменных без статичного текста письма, а имена переменных не совпадали с теми,
что отдаёт services/values.py. В итоге в конструкторе и в Word не было ни шапки,
ни обращений, ни подписей.

Сборка идёт по тексту подлинного письма (data/templates/<code>/example_filled.real.pdf):
шапка агентства, номер, тема, адресат, два блока данных (подписант и клиент),
тело, подпись. Значения — переменные в {{ }} с именами из values.py.

Запуск:
    /usr/bin/python3 scripts/build_bridging_letter_docx.py            # собрать
    /usr/bin/python3 scripts/build_bridging_letter_docx.py --render   # собрать и отрендерить пробный PDF
"""
from __future__ import annotations

import shutil
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

ROOT = Path('/opt/wolf_v2')
CODE = 'bridging_request_letter__default'
TPL = ROOT / 'data' / 'templates' / CODE / 'word' / 'template.docx'

LETTERHEAD = [
    'PT ALPHA PRIME BALI',
    'Jl. Wisma Udayana No.164 Buruan, Kec. Blahbatuh, Kabupaten Gianyar, Bali 80581',
    'Telp: +62 811 397 9669, Email: ptalphaprimeconsulting@gmail.com',
]


def _p(doc, text: str = '', size: float = 11, bold: bool = False,
       align=None, space_before: float = 0, space_after: float = 0,
       style: str | None = None):
    par = doc.add_paragraph(style=style)
    if align is not None:
        par.alignment = align
    par.paragraph_format.space_before = Pt(space_before)
    par.paragraph_format.space_after = Pt(space_after)
    par.paragraph_format.line_spacing = 1.15
    if text:
        run = par.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
        run.font.name = 'Times New Roman'
    return par


def _field(doc, label: str, value: str, size: float = 11) -> None:
    """Строка «Метка : значение» — метка у левого края, двоеточие и значение по табуляции."""
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.15
    pf.tab_stops.add_tab_stop(Cm(4.2))
    pf.tab_stops.add_tab_stop(Cm(4.6))
    r1 = par.add_run(f'{label}\t:\t{value}')
    r1.font.size = Pt(size)
    r1.font.name = 'Times New Roman'


def build() -> Path:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(2.2), Cm(2.0)
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)

    # Шапка агентства
    _p(doc, LETTERHEAD[0], size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, LETTERHEAD[1], size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, LETTERHEAD[2], size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    # Линия под шапкой
    line = doc.add_paragraph()
    line.paragraph_format.space_before = Pt(0)
    line.paragraph_format.space_after = Pt(8)
    pPr = line._p.get_or_add_pPr()
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '000000')
    pbdr.append(bottom)
    pPr.append(pbdr)

    # Реквизиты письма
    _p(doc, 'Nomor: {{ letter_no }}')
    _p(doc, 'Lampiran: satu berkas')
    _p(doc, 'Hal: {{ letter_subject }}', space_after=10)

    # Адресат
    _p(doc, 'Kepada Yth.')
    _p(doc, '{{ addressee_line1 }}')
    _p(doc, '{{ addressee_line2 }}')
    _p(doc, '{{ addressee_line3 }}', space_after=10)

    # Подписант (доверенное лицо агентства)
    _p(doc, 'Bersama ini dengan hormat kami:', space_after=2)
    _field(doc, 'Nama', '{{ attorney_name }}')
    _field(doc, 'Pekerjaan', 'Swasta')
    _field(doc, 'Kewarganegaraan', '{{ attorney_citizenship }}')
    _field(doc, 'NIK', '{{ attorney_nik }}')
    _field(doc, 'Alamat', '{{ attorney_address }}')

    # Переход к данным клиента
    _p(doc, 'selaku penerima kuasa dari {{ sponsor_company }}, dengan ini kami mengajukan '
            'permohonan', space_before=8)
    _p(doc, '{{ body_request_line }}', space_after=2)

    # Данные клиента
    _field(doc, 'Nama', '{{ client_name }}')
    _field(doc, 'Jenis Kelamin', '{{ client_gender }}')
    _field(doc, 'Tempat/Tgl. Lahir', '{{ place_birth_inline }}')
    _field(doc, 'Kewarganegaraan', '{{ citizenship }}')
    _field(doc, 'No. Paspor', '{{ passport_no }}')
    _field(doc, 'Masa Berlaku Paspor', '{{ passport_expiry }}')
    _field(doc, 'Nomor ITK', '{{ itk_no }}')
    _field(doc, 'Masa Berlaku ITK', '{{ itk_expiry }}')
    _field(doc, 'Alamat', '{{ client_address }}')

    # Закрытие и подпись
    _p(doc, 'Demikian permohonan yang dapat kami ajukan, atas perhatiannya kami ucapkan '
            'terima kasih.', space_before=10, space_after=14)
    _p(doc, '{{ place_date }}', align=WD_ALIGN_PARAGRAPH.RIGHT)
    _p(doc, 'Hormat kami,', align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=30)
    _p(doc, '{{ attorney_name }}', align=WD_ALIGN_PARAGRAPH.RIGHT)

    TPL.parent.mkdir(parents=True, exist_ok=True)
    if TPL.exists():
        bak = TPL.with_name(f'template.docx.bak_{datetime.now():%Y%m%d_%H%M}')
        shutil.copy2(TPL, bak)
        print('бэкап:', bak.name)
    doc.save(TPL)
    print('собрано:', TPL, TPL.stat().st_size, 'байт')
    return TPL


def render() -> None:
    """Пробный PDF с реальными значениями (для проверки глазами)."""
    import docxtpl
    import sqlite3
    sys.path.insert(0, str(ROOT))
    from services.values import build_template_values

    c = sqlite3.connect(ROOT / 'data' / 'wolf.db')
    c.row_factory = sqlite3.Row
    app = 109
    a = dict(c.execute('SELECT * FROM applications WHERE id=?', (app,)).fetchone())
    cl = dict(c.execute('SELECT * FROM clients WHERE id=?', (a['client_id'],)).fetchone())
    vals = build_template_values(cl, CODE, {}, app) or {}
    out = ROOT / 'data' / 'renders' / f'{CODE}_check.docx'
    out.parent.mkdir(parents=True, exist_ok=True)
    tpl_obj = docxtpl.DocxTemplate(str(TPL))
    tpl_obj.render({k: ('' if v is None else str(v)) for k, v in vals.items()})
    tpl_obj.save(str(out))
    print('DOCX:', out)
    import subprocess
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', str(out),
                    '--outdir', str(out.parent)], check=True, stdout=subprocess.DEVNULL)
    pdf = out.with_suffix('.pdf')
    print('PDF:', pdf)
    import pymupdf
    d = pymupdf.open(pdf)
    pix = d[0].get_pixmap(dpi=110)
    png = ROOT / 'data' / 'renders' / f'{CODE}_check.png'
    pix.save(png)
    print('PNG:', png)


if __name__ == '__main__':
    build()
    if '--render' in sys.argv:
        render()
